"""
Word Document Extractor — pulls text from .docx files using python-docx.
Preserves heading structure as markdown-style headers.
"""

from pathlib import Path

MAX_CHARS = 40000


def extract_docx(file_path: Path) -> str:
    """Extract text from a .docx file. Returns plain text with heading markers."""
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(file_path)
    lines = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        style = para.style.name if para.style else ""

        # Map Word heading styles to markdown headers
        if style.startswith("Heading 1"):
            lines.append(f"# {text}")
        elif style.startswith("Heading 2"):
            lines.append(f"## {text}")
        elif style.startswith("Heading 3"):
            lines.append(f"### {text}")
        else:
            lines.append(text)

    # Also pull text from tables
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(cells))

    if not lines:
        return "[Word document contained no extractable text]"

    full_text = "\n\n".join(lines)
    if len(full_text) > MAX_CHARS:
        full_text = full_text[:MAX_CHARS] + "\n\n[... truncated — document exceeded limit ...]"

    return full_text
